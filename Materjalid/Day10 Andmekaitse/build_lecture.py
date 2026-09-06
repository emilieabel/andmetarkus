# -*- coding: utf-8 -*-
"""Andmetarkus · Päev 10: Andmekaitse — 3-tunnine loeng (slaidid + notes)."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
from pathlib import Path

# --- kujundus (kollane bänner + must tekst) ---
YELLOW = RGBColor(255, 210, 0)
YELLOW_DARK = RGBColor(230, 180, 0)
BLACK = RGBColor(18, 18, 18)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(32, 32, 32)
GRAY = RGBColor(88, 88, 88)
CREAM = RGBColor(255, 248, 214)
SOFT = RGBColor(250, 250, 247)
MUTED = RGBColor(110, 110, 110)

W, H = Inches(13.333), Inches(7.5)
FONT = "Calibri"
BANNER = Path(__file__).parent / "teema_banner.png"
OUT = Path(__file__).parent / "Andmekaitse_loeng_Andmetarkus.pptx"


def _set_run(run, text, size, bold=False, color=BLACK, italic=False, name=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def add_textbox(slide, l, t, w, h, text, size=20, bold=False, color=BLACK,
                align=PP_ALIGN.LEFT, italic=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run(), text, size, bold, color, italic)
    return tf


def add_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def set_notes(slide, text):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    tf.text = text
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.size = Pt(14)
            run.font.name = FONT


def footer(slide, num, total, dark=False):
    c = WHITE if dark else MUTED
    add_textbox(slide, Inches(0.4), Inches(7.15), Inches(10.5), Inches(0.28),
                "Andmetarkus  ·  Päev 10: Andmekaitse", 11, False, c)
    add_textbox(slide, Inches(11.2), Inches(7.15), Inches(1.7), Inches(0.28),
                f"{num} / {total}", 11, False, c, PP_ALIGN.RIGHT)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bullets_box(slide, l, t, w, h, items, size=22, color=CHARCOAL, spacing=10):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(spacing)
        p.alignment = PP_ALIGN.LEFT
        if isinstance(item, tuple):
            text, lvl = item
            p.level = lvl
        else:
            text = item
        prefix = "•  " if p.level == 0 else "–  "
        _set_run(p.add_run(), prefix + text, size if p.level == 0 else size - 2,
                 False, color if p.level == 0 else GRAY)
    return tf


def card(slide, l, t, w, h, title, body, title_size=16, body_size=14):
    add_round(slide, l, t, w, h, CREAM)
    add_textbox(slide, l + Inches(0.18), t + Inches(0.14), w - Inches(0.36), Inches(0.4),
                title, title_size, True, BLACK)
    add_textbox(slide, l + Inches(0.18), t + Inches(0.52), w - Inches(0.36), h - Inches(0.66),
                body, body_size, False, CHARCOAL)


# ---------- slaiditüübid ----------

def slide_title(prs, spec, num, total):
    s = blank(prs)
    add_rect(s, 0, 0, W, H, YELLOW)
    add_rect(s, 0, 0, Inches(0.22), H, BLACK)
    add_textbox(s, Inches(0.7), Inches(1.55), Inches(12), Inches(0.4),
                spec.get("kicker", "ANDMETARKUS  ·  PÄEV 10"), 16, True, BLACK)
    add_textbox(s, Inches(0.7), Inches(2.05), Inches(12), Inches(1.6),
                spec["title"], 40, True, BLACK)
    add_textbox(s, Inches(0.7), Inches(3.85), Inches(12), Inches(1.2),
                spec.get("subtitle", ""), 20, False, CHARCOAL)
    add_textbox(s, Inches(0.7), Inches(5.4), Inches(12), Inches(0.8),
                spec.get("meta", "3-tunnine loeng algajatele  ·  Eesti õigusruum: IKÜM + IKS"), 16, False, BLACK)
    set_notes(s, spec["notes"])
    return s


def slide_section(prs, spec, num, total):
    s = blank(prs)
    add_rect(s, 0, 0, W, H, YELLOW)
    add_rect(s, 0, 0, Inches(0.22), H, BLACK)
    add_textbox(s, Inches(0.7), Inches(1.7), Inches(12), Inches(0.4),
                spec.get("kicker", ""), 16, True, BLACK)
    add_textbox(s, Inches(0.7), Inches(2.2), Inches(12), Inches(2.2),
                spec["title"], 36, True, BLACK)
    add_textbox(s, Inches(0.7), Inches(4.6), Inches(12), Inches(1.2),
                spec.get("subtitle", ""), 18, False, CHARCOAL)
    footer(s, num, total, dark=False)
    set_notes(s, spec["notes"])
    return s


def slide_content(prs, spec, num, total):
    s = blank(prs)
    add_rect(s, 0, 0, W, H, WHITE)
    add_rect(s, 0, 0, W, Inches(0.12), YELLOW)
    add_rect(s, 0, 0, Inches(0.14), H, YELLOW)
    if spec.get("kicker"):
        add_textbox(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.32),
                    spec["kicker"], 13, True, MUTED)
        title_top = Inches(0.52)
    else:
        title_top = Inches(0.32)
    add_textbox(s, Inches(0.55), title_top, Inches(12.2), Inches(0.85),
                spec["title"], 28, True, BLACK)
    items = spec.get("bullets", [])
    has_callout = bool(spec.get("callout"))
    bullet_h = Inches(4.55) if has_callout else Inches(5.4)
    if items:
        bullets_box(s, Inches(0.55), Inches(1.45), Inches(12.2), bullet_h,
                    items, spec.get("size", 22), CHARCOAL, spec.get("spacing", 12))
    if has_callout:
        add_round(s, Inches(0.55), Inches(6.15), Inches(12.2), Inches(0.85), CREAM)
        add_textbox(s, Inches(0.75), Inches(6.28), Inches(11.8), Inches(0.6),
                    spec["callout"], 16, True, BLACK)
    footer(s, num, total)
    set_notes(s, spec["notes"])
    return s


def slide_cards(prs, spec, num, total):
    s = blank(prs)
    add_rect(s, 0, 0, W, H, WHITE)
    add_rect(s, 0, 0, W, Inches(0.12), YELLOW)
    add_rect(s, 0, 0, Inches(0.14), H, YELLOW)
    add_textbox(s, Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.7),
                spec["title"], 28, True, BLACK)
    cards = spec["cards"]
    n = len(cards)
    gap = Inches(0.22)
    left = Inches(0.55)
    usable = Inches(12.2)
    cw = (usable - gap * (n - 1)) / n
    top = Inches(1.3)
    ch = Inches(5.4) if not spec.get("subtitle") else Inches(5.15)
    if spec.get("subtitle"):
        add_textbox(s, Inches(0.55), Inches(1.05), Inches(12.2), Inches(0.35),
                    spec["subtitle"], 16, False, GRAY)
        top = Inches(1.45)
        ch = Inches(5.25)
    for i, c in enumerate(cards):
        x = left + i * (cw + gap)
        card(s, x, top, cw, ch, c[0], c[1], 16, 14)
    footer(s, num, total)
    set_notes(s, spec["notes"])
    return s


def slide_quote(prs, spec, num, total):
    s = blank(prs)
    add_rect(s, 0, 0, W, H, CHARCOAL)
    add_rect(s, 0, 0, Inches(0.22), H, YELLOW)
    add_textbox(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(2.8),
                spec["title"], 28, True, WHITE)
    add_textbox(s, Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.0),
                spec.get("subtitle", ""), 18, False, YELLOW)
    footer(s, num, total, dark=True)
    set_notes(s, spec["notes"])
    return s


def slide_exercise(prs, spec, num, total):
    s = blank(prs)
    add_rect(s, 0, 0, W, H, WHITE)
    add_rect(s, 0, 0, W, Inches(0.12), YELLOW)
    add_rect(s, 0, 0, Inches(0.14), H, YELLOW)
    add_round(s, Inches(0.55), Inches(0.28), Inches(2.1), Inches(0.4), YELLOW)
    add_textbox(s, Inches(0.55), Inches(0.30), Inches(2.1), Inches(0.36),
                spec.get("badge", "HARJUTUS"), 13, True, BLACK, PP_ALIGN.CENTER)
    add_textbox(s, Inches(2.8), Inches(0.28), Inches(9.8), Inches(0.4),
                spec.get("time", "6–8 minutit  ·  paarides või väikestes gruppides"), 14, False, GRAY)
    add_textbox(s, Inches(0.55), Inches(0.8), Inches(12.2), Inches(0.7),
                spec["title"], 26, True, BLACK)
    bullets_box(s, Inches(0.55), Inches(1.6), Inches(12.2), Inches(5.2),
                spec.get("bullets", []), 20, CHARCOAL, 10)
    footer(s, num, total)
    set_notes(s, spec["notes"])
    return s


def slide_break(prs, spec, num, total):
    s = blank(prs)
    add_rect(s, 0, 0, W, H, YELLOW)
    add_rect(s, 0, 0, Inches(0.22), H, BLACK)
    add_textbox(s, Inches(0.7), Inches(2.4), Inches(12), Inches(1.2),
                spec["title"], 44, True, BLACK)
    add_textbox(s, Inches(0.7), Inches(3.8), Inches(12), Inches(1.0),
                spec.get("subtitle", "10 minutit  ·  oleme tagasi täpselt ajaks"), 20, False, CHARCOAL)
    footer(s, num, total)
    set_notes(s, spec["notes"])
    return s


RENDER = {
    "title": slide_title,
    "section": slide_section,
    "content": slide_content,
    "cards": slide_cards,
    "quote": slide_quote,
    "exercise": slide_exercise,
    "break": slide_break,
}


def build(slides):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = len(slides)
    for i, spec in enumerate(slides, 1):
        RENDER[spec["type"]](prs, spec, i, total)
    prs.save(OUT)
    return OUT, total
